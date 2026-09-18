"""
Phishing Awareness Presentation Generator
==========================================
Generates a 13-slide .pptx with:
  - Dark cybersecurity theme (charcoal/navy + cyan accents)
  - Educational slides (1-7)
  - Interactive "Spot the Phish" game (8-13)
  - Full speaker notes from the script
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ── Theme Colors ──────────────────────────────────────────────────────────
BG_DARK       = RGBColor(0x1A, 0x1A, 0x2E)  # Deep navy-charcoal
BG_CARD       = RGBColor(0x16, 0x21, 0x3E)  # Slightly lighter card bg
BG_EMAIL      = RGBColor(0x25, 0x25, 0x3A)  # Email mockup bg
WHITE         = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY    = RGBColor(0xCC, 0xCC, 0xCC)
CYAN          = RGBColor(0x00, 0xD4, 0xFF)
RED_ACCENT    = RGBColor(0xFF, 0x44, 0x44)
GREEN_ACCENT  = RGBColor(0x44, 0xFF, 0x88)
YELLOW_ACCENT = RGBColor(0xFF, 0xD7, 0x00)
ORANGE_ACCENT = RGBColor(0xFF, 0x8C, 0x00)
DIM_GRAY      = RGBColor(0x88, 0x88, 0x99)
SMS_GREEN     = RGBColor(0x25, 0xD3, 0x66)
SMS_BUBBLE    = RGBColor(0x30, 0x30, 0x48)
BUTTON_BLUE   = RGBColor(0x00, 0x6A, 0xB6)


def set_slide_bg(slide, color):
    """Set solid background color for a slide."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height):
    """Add a textbox and return the text frame."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    return tf


def set_paragraph(tf, text, font_size=18, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT,
                  font_name="Calibri", space_after=Pt(6), space_before=Pt(0), italic=False):
    """Configure the first (default) paragraph of a text frame."""
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.font.italic = italic
    p.alignment = alignment
    p.space_after = space_after
    p.space_before = space_before
    return p


def add_paragraph(tf, text, font_size=18, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT,
                  font_name="Calibri", space_after=Pt(6), space_before=Pt(0), italic=False,
                  level=0):
    """Add a new paragraph to an existing text frame."""
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.font.italic = italic
    p.alignment = alignment
    p.space_after = space_after
    p.space_before = space_before
    p.level = level
    return p


def add_rounded_rect(slide, left, top, width, height, fill_color, border_color=None):
    """Add a rounded rectangle shape."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape


def add_rectangle(slide, left, top, width, height, fill_color, border_color=None):
    """Add a plain rectangle shape."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape


def add_oval(slide, left, top, width, height, border_color=RED_ACCENT, fill=False):
    """Add an oval shape (used for red circle annotations)."""
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, width, height)
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = border_color
    else:
        shape.fill.background()
    shape.line.color.rgb = border_color
    shape.line.width = Pt(3)
    return shape


def add_speaker_notes(slide, text):
    """Add speaker notes to a slide."""
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = text


def add_icon_circle(slide, left, top, text, bg_color=CYAN, text_color=BG_DARK, size=Inches(0.6)):
    """Add a small circle with an icon character or number."""
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(16)
    p.font.color.rgb = text_color
    p.font.bold = True
    p.font.name = "Calibri"
    p.alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].space_before = Pt(0)
    tf.paragraphs[0].space_after = Pt(0)


def add_red_flag_annotation(slide, left, top, width, text, icon="⚠"):
    """Add a red flag annotation box with icon."""
    # Background pill
    box = add_rounded_rect(slide, left, top, width, Inches(0.55), RED_ACCENT)
    box_tf = box.text_frame
    box_tf.word_wrap = True
    p = box_tf.paragraphs[0]
    p.text = f"{icon} {text}"
    p.font.size = Pt(13)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.font.name = "Calibri"
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    return box


# ══════════════════════════════════════════════════════════════════════════
# SLIDE BUILDERS
# ══════════════════════════════════════════════════════════════════════════

def build_slide_1(prs):
    """Title Slide - Hook & Introduction"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_slide_bg(slide, BG_DARK)

    # Decorative top accent bar
    add_rectangle(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.06), CYAN)

    # Shield icon emoji
    tf = add_textbox(slide, Inches(0), Inches(1.5), Inches(13.33), Inches(1.2))
    set_paragraph(tf, "🛡️", font_size=60, alignment=PP_ALIGN.CENTER)

    # Title
    tf = add_textbox(slide, Inches(1.5), Inches(2.6), Inches(10.33), Inches(1.2))
    set_paragraph(tf, "PHISHING AWARENESS", font_size=44, color=WHITE, bold=True,
                  alignment=PP_ALIGN.CENTER, font_name="Calibri")
    add_paragraph(tf, "Don't Take the Bait", font_size=24, color=CYAN,
                  alignment=PP_ALIGN.CENTER, space_before=Pt(8))

    # Subtitle
    tf = add_textbox(slide, Inches(2.5), Inches(4.2), Inches(8.33), Inches(0.8))
    set_paragraph(tf, "How to identify and defend against phishing attacks",
                  font_size=18, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)
    add_paragraph(tf, "Presented by Anonymous SDMCET",
                  font_size=14, color=DIM_GRAY, alignment=PP_ALIGN.CENTER, space_before=Pt(6))

    # Bottom decorative bar
    add_rectangle(slide, Inches(0), Inches(7.44), Inches(13.33), Inches(0.06), CYAN)

    # Club branding
    tf = add_textbox(slide, Inches(2.5), Inches(5.2), Inches(8.33), Inches(0.6))
    set_paragraph(tf, "ANONYMOUS SDMCET",
                  font_size=22, color=CYAN, bold=True, alignment=PP_ALIGN.CENTER,
                  font_name="Calibri")

    # Bottom info
    tf = add_textbox(slide, Inches(3), Inches(6.5), Inches(7.33), Inches(0.5))
    set_paragraph(tf, "Cybersecurity Club  |  Interactive Training Session",
                  font_size=14, color=DIM_GRAY, alignment=PP_ALIGN.CENTER)

    add_speaker_notes(slide, (
        "Welcome, everyone. Today, we are going to talk about the digital equivalent of "
        "a con artist knocking on your front door: Phishing. Cybercriminals send billions "
        "of phishing emails every day, and their entire goal is to manipulate you into "
        "handing over your sensitive information, your passwords, or your money. By the "
        "end of this session, you will know exactly how to spot them."
    ))


def build_slide_2(prs):
    """What is Phishing?"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)
    add_rectangle(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.06), CYAN)

    # Title
    tf = add_textbox(slide, Inches(0.6), Inches(0.3), Inches(12), Inches(0.8))
    set_paragraph(tf, "🎣  What is Phishing?", font_size=36, color=CYAN, bold=True)

    # Definition card
    card = add_rounded_rect(slide, Inches(0.8), Inches(1.4), Inches(11.73), Inches(1.5), BG_CARD, CYAN)
    card_tf = card.text_frame
    card_tf.word_wrap = True
    set_paragraph(card_tf, "A social engineering attack where cybercriminals impersonate trusted "
                  "entities to trick you into revealing sensitive information.",
                  font_size=20, color=WHITE, alignment=PP_ALIGN.CENTER)

    # Three columns: They pretend to be | They target | They want
    col_width = Inches(3.5)
    col_starts = [Inches(0.8), Inches(4.9), Inches(9.0)]
    col_tops = Inches(3.4)
    col_height = Inches(3.5)

    titles = ["🎭  They Pretend To Be...", "🧠  They Target...", "💰  They Want..."]
    items = [
        ["Your bank", "Your boss / CEO", "Netflix, Amazon, FedEx", "Tech support", "Government agencies"],
        ["Your trust", "Your emotions (fear, greed)", "Your urgency to act", "Your lack of attention", "Your helpful nature"],
        ["Passwords & logins", "Credit card numbers", "Personal data (SSN, DOB)", "Money transfers", "Access to your systems"],
    ]

    for idx in range(3):
        col_card = add_rounded_rect(slide, col_starts[idx], col_tops, col_width, col_height, BG_CARD)
        col_tf = col_card.text_frame
        col_tf.word_wrap = True
        set_paragraph(col_tf, titles[idx], font_size=17, color=CYAN, bold=True,
                      alignment=PP_ALIGN.CENTER, space_after=Pt(14))
        for item in items[idx]:
            add_paragraph(col_tf, f"•  {item}", font_size=15, color=LIGHT_GRAY,
                          alignment=PP_ALIGN.LEFT, space_after=Pt(8))

    # Bottom tagline
    tf = add_textbox(slide, Inches(1), Inches(7.0), Inches(11.33), Inches(0.4))
    set_paragraph(tf, "\"They aren't hacking your computer — they are hacking your psychology.\"",
                  font_size=16, color=YELLOW_ACCENT, alignment=PP_ALIGN.CENTER, italic=True)

    add_speaker_notes(slide, (
        "Phishing is a social engineering attack. Attackers disguise themselves as a trusted "
        "entity—like your bank, your boss, or a service you use like Netflix or Amazon. They "
        "aren't hacking your computer; they are hacking your psychology. They want you to make "
        "a quick, emotional decision without thinking."
    ))


def build_slide_3(prs):
    """Red Flag 1 - The Sense of Urgency"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)
    add_rectangle(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.06), RED_ACCENT)

    # Tag
    tag = add_rounded_rect(slide, Inches(0.6), Inches(0.3), Inches(1.8), Inches(0.45), RED_ACCENT)
    tag_tf = tag.text_frame
    set_paragraph(tag_tf, "RED FLAG #1", font_size=14, color=WHITE, bold=True,
                  alignment=PP_ALIGN.CENTER)

    # Title
    tf = add_textbox(slide, Inches(2.7), Inches(0.25), Inches(10), Inches(0.7))
    set_paragraph(tf, "⏰  The Sense of Urgency", font_size=34, color=WHITE, bold=True)

    # Main content card
    card = add_rounded_rect(slide, Inches(0.8), Inches(1.3), Inches(11.73), Inches(2.2), BG_CARD, RED_ACCENT)
    card_tf = card.text_frame
    card_tf.word_wrap = True
    set_paragraph(card_tf, "Phishers want you to PANIC and ACT without THINKING.",
                  font_size=24, color=RED_ACCENT, bold=True, alignment=PP_ALIGN.CENTER,
                  space_after=Pt(16))
    add_paragraph(card_tf, "If you feel RUSHED, you are likely being PHISHED.",
                  font_size=20, color=YELLOW_ACCENT, alignment=PP_ALIGN.CENTER, bold=True)

    # Example phrases
    examples_top = Inches(3.9)
    phrases = [
        ("⚡", "\"Your account will be\nsuspended in 24 hours!\""),
        ("⚖️", "\"You are facing\nimmediate legal action!\""),
        ("🎰", "\"You've won! Claim your\nprize before it expires!\""),
    ]
    for i, (icon, phrase) in enumerate(phrases):
        x = Inches(0.8) + i * Inches(4.0)
        box = add_rounded_rect(slide, x, examples_top, Inches(3.7), Inches(1.7), BG_EMAIL, RED_ACCENT)
        box_tf = box.text_frame
        box_tf.word_wrap = True
        set_paragraph(box_tf, icon, font_size=30, alignment=PP_ALIGN.CENTER, space_after=Pt(4))
        add_paragraph(box_tf, phrase, font_size=16, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    # Bottom advice
    advice_card = add_rounded_rect(slide, Inches(1.5), Inches(6.0), Inches(10.33), Inches(0.9),
                                   BG_CARD, GREEN_ACCENT)
    advice_tf = advice_card.text_frame
    advice_tf.word_wrap = True
    set_paragraph(advice_tf, "✅  Real organizations rarely demand immediate, panicked action. "
                  "When in doubt, contact the company directly using a known phone number or website.",
                  font_size=16, color=GREEN_ACCENT, alignment=PP_ALIGN.CENTER)

    add_speaker_notes(slide, (
        "The number one tool of a phisher is panic. If a message says your account will be "
        "suspended in 24 hours, or you're facing legal action, or you just won a lottery you "
        "didn't enter—stop. Real organizations rarely demand immediate, panicked action. If you "
        "feel rushed, you are likely being phished."
    ))


def build_slide_4(prs):
    """Red Flag 2 - The Sender Address"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)
    add_rectangle(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.06), RED_ACCENT)

    tag = add_rounded_rect(slide, Inches(0.6), Inches(0.3), Inches(1.8), Inches(0.45), RED_ACCENT)
    tag_tf = tag.text_frame
    set_paragraph(tag_tf, "RED FLAG #2", font_size=14, color=WHITE, bold=True,
                  alignment=PP_ALIGN.CENTER)

    tf = add_textbox(slide, Inches(2.7), Inches(0.25), Inches(10), Inches(0.7))
    set_paragraph(tf, "📧  The Sender Address", font_size=34, color=WHITE, bold=True)

    # Display name vs real address comparison
    card = add_rounded_rect(slide, Inches(0.8), Inches(1.3), Inches(11.73), Inches(1.4), BG_CARD)
    card_tf = card.text_frame
    card_tf.word_wrap = True
    set_paragraph(card_tf, "Never trust the display name. Always check the ACTUAL email address.",
                  font_size=22, color=YELLOW_ACCENT, bold=True, alignment=PP_ALIGN.CENTER,
                  space_after=Pt(12))
    add_paragraph(card_tf, "The display name can say anything. The domain after @ is what matters.",
                  font_size=17, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    # Fake vs Real examples - side by side
    # Left: Fake
    fake_card = add_rounded_rect(slide, Inches(0.8), Inches(3.1), Inches(5.6), Inches(3.4),
                                 BG_EMAIL, RED_ACCENT)
    fake_tf = fake_card.text_frame
    fake_tf.word_wrap = True
    set_paragraph(fake_tf, "❌  FAKE", font_size=20, color=RED_ACCENT, bold=True,
                  alignment=PP_ALIGN.CENTER, space_after=Pt(12))
    add_paragraph(fake_tf, "Display:  PayPal Support", font_size=15, color=LIGHT_GRAY,
                  space_after=Pt(4))
    add_paragraph(fake_tf, "Actual:   support@paypa1-update-security.com", font_size=14,
                  color=RED_ACCENT, bold=True, space_after=Pt(16))
    add_paragraph(fake_tf, "Display:  Apple ID", font_size=15, color=LIGHT_GRAY,
                  space_after=Pt(4))
    add_paragraph(fake_tf, "Actual:   noreply@apple-id-verify.info", font_size=14,
                  color=RED_ACCENT, bold=True, space_after=Pt(16))
    add_paragraph(fake_tf, "Display:  Microsoft Account", font_size=15, color=LIGHT_GRAY,
                  space_after=Pt(4))
    add_paragraph(fake_tf, "Actual:   alert@microsoft365-security.org", font_size=14,
                  color=RED_ACCENT, bold=True)

    # Right: Real
    real_card = add_rounded_rect(slide, Inches(6.93), Inches(3.1), Inches(5.6), Inches(3.4),
                                 BG_EMAIL, GREEN_ACCENT)
    real_tf = real_card.text_frame
    real_tf.word_wrap = True
    set_paragraph(real_tf, "✅  LEGITIMATE", font_size=20, color=GREEN_ACCENT, bold=True,
                  alignment=PP_ALIGN.CENTER, space_after=Pt(12))
    add_paragraph(real_tf, "Display:  PayPal", font_size=15, color=LIGHT_GRAY,
                  space_after=Pt(4))
    add_paragraph(real_tf, "Actual:   service@paypal.com", font_size=14,
                  color=GREEN_ACCENT, bold=True, space_after=Pt(16))
    add_paragraph(real_tf, "Display:  Apple", font_size=15, color=LIGHT_GRAY,
                  space_after=Pt(4))
    add_paragraph(real_tf, "Actual:   noreply@apple.com", font_size=14,
                  color=GREEN_ACCENT, bold=True, space_after=Pt(16))
    add_paragraph(real_tf, "Display:  Microsoft Account Team", font_size=15, color=LIGHT_GRAY,
                  space_after=Pt(4))
    add_paragraph(real_tf, "Actual:   account-security@microsoft.com", font_size=14,
                  color=GREEN_ACCENT, bold=True)

    # Bottom tip
    tip = add_rounded_rect(slide, Inches(1.5), Inches(6.8), Inches(10.33), Inches(0.6),
                           BG_CARD, CYAN)
    tip_tf = tip.text_frame
    tip_tf.word_wrap = True
    set_paragraph(tip_tf, "💡  Watch for: misspellings (paypa1 vs paypal), extra subdomains, "
                  "public domains (@gmail.com), unusual TLDs (.info, .xyz)",
                  font_size=14, color=CYAN, alignment=PP_ALIGN.CENTER)

    add_speaker_notes(slide, (
        "Never trust the display name. An email might say it's from 'PayPal Support,' but if "
        "you look at the actual email address, it might say something like "
        "'support@paypa1-update-security.com'. Always verify the domain name after the '@' "
        "symbol. If it looks overly complicated, misspelled, or uses a public domain like "
        "@gmail.com for an official business, it's a trap."
    ))


def build_slide_5(prs):
    """Red Flag 3 - Suspicious Links and Attachments"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)
    add_rectangle(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.06), RED_ACCENT)

    tag = add_rounded_rect(slide, Inches(0.6), Inches(0.3), Inches(1.8), Inches(0.45), RED_ACCENT)
    tag_tf = tag.text_frame
    set_paragraph(tag_tf, "RED FLAG #3", font_size=14, color=WHITE, bold=True,
                  alignment=PP_ALIGN.CENTER)

    tf = add_textbox(slide, Inches(2.7), Inches(0.25), Inches(10), Inches(0.7))
    set_paragraph(tf, "🔗  Suspicious Links & Attachments", font_size=34, color=WHITE, bold=True)

    # --- Links section ---
    link_card = add_rounded_rect(slide, Inches(0.8), Inches(1.3), Inches(5.6), Inches(5.2),
                                 BG_CARD, ORANGE_ACCENT)
    link_tf = link_card.text_frame
    link_tf.word_wrap = True
    set_paragraph(link_tf, "🖱️  HOVER BEFORE YOU CLICK", font_size=20, color=ORANGE_ACCENT,
                  bold=True, alignment=PP_ALIGN.CENTER, space_after=Pt(14))
    add_paragraph(link_tf, "When you hover over a link, a tooltip shows the real URL.",
                  font_size=15, color=LIGHT_GRAY, space_after=Pt(16))
    add_paragraph(link_tf, "What you SEE:", font_size=14, color=DIM_GRAY, bold=True,
                  space_after=Pt(4))
    add_paragraph(link_tf, "  \"Click here to update your billing\"", font_size=16,
                  color=CYAN, space_after=Pt(12))
    add_paragraph(link_tf, "What it ACTUALLY links to:", font_size=14, color=DIM_GRAY,
                  bold=True, space_after=Pt(4))
    add_paragraph(link_tf, "  http://x7k2-billing.sketchy-site.ru/steal",
                  font_size=14, color=RED_ACCENT, bold=True, space_after=Pt(16))
    add_paragraph(link_tf, "⚠ Watch for:", font_size=15, color=YELLOW_ACCENT, bold=True,
                  space_after=Pt(6))
    add_paragraph(link_tf, "• Shortened URLs (bit.ly, tinyurl)", font_size=14,
                  color=LIGHT_GRAY, space_after=Pt(4))
    add_paragraph(link_tf, "• Misspelled domains", font_size=14,
                  color=LIGHT_GRAY, space_after=Pt(4))
    add_paragraph(link_tf, "• HTTP instead of HTTPS", font_size=14,
                  color=LIGHT_GRAY, space_after=Pt(4))
    add_paragraph(link_tf, "• Random subdomains", font_size=14,
                  color=LIGHT_GRAY)

    # --- Attachments section ---
    attach_card = add_rounded_rect(slide, Inches(6.93), Inches(1.3), Inches(5.6), Inches(5.2),
                                   BG_CARD, RED_ACCENT)
    attach_tf = attach_card.text_frame
    attach_tf.word_wrap = True
    set_paragraph(attach_tf, "📎  NEVER OPEN UNEXPECTED FILES", font_size=20, color=RED_ACCENT,
                  bold=True, alignment=PP_ALIGN.CENTER, space_after=Pt(14))
    add_paragraph(attach_tf, "Common malicious attachment types:", font_size=15,
                  color=LIGHT_GRAY, space_after=Pt(12))

    danger_files = [
        ("Invoice_2024.exe", "Executable — runs malware directly"),
        ("Report.pdf.scr", "Double extension — hides true type"),
        ("Tracking_Info.zip", "Archive — may contain malware"),
        ("Document.docm", "Macro-enabled — can run code"),
        ("Payment.html", "Can redirect to phishing page"),
    ]
    for fname, desc in danger_files:
        add_paragraph(attach_tf, f"❌  {fname}", font_size=15, color=RED_ACCENT, bold=True,
                      space_after=Pt(2))
        add_paragraph(attach_tf, f"     {desc}", font_size=13, color=DIM_GRAY,
                      space_after=Pt(10))

    # Bottom advice
    tip = add_rounded_rect(slide, Inches(1.5), Inches(6.8), Inches(10.33), Inches(0.6),
                           BG_CARD, GREEN_ACCENT)
    tip_tf = tip.text_frame
    tip_tf.word_wrap = True
    set_paragraph(tip_tf, "✅  Rule of thumb: If you didn't expect it, don't click it. "
                  "Verify with the sender through a separate channel.",
                  font_size=15, color=GREEN_ACCENT, alignment=PP_ALIGN.CENTER)

    add_speaker_notes(slide, (
        "Before you click any link, hover your mouse over it. A small box will appear showing "
        "you the actual destination URL. If the button says 'Update Billing' but the URL goes "
        "to a random, unrelated website, do not click it. Similarly, never download unexpected "
        "attachments, especially unexpected invoices or receipts—these are classic delivery "
        "methods for malware."
    ))


def build_slide_6(prs):
    """Red Flag 4 - Generic Greetings"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)
    add_rectangle(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.06), RED_ACCENT)

    tag = add_rounded_rect(slide, Inches(0.6), Inches(0.3), Inches(1.8), Inches(0.45), RED_ACCENT)
    tag_tf = tag.text_frame
    set_paragraph(tag_tf, "RED FLAG #4", font_size=14, color=WHITE, bold=True,
                  alignment=PP_ALIGN.CENTER)

    tf = add_textbox(slide, Inches(2.7), Inches(0.25), Inches(10), Inches(0.7))
    set_paragraph(tf, "👤  Generic Greetings", font_size=34, color=WHITE, bold=True)

    # Main point card
    card = add_rounded_rect(slide, Inches(0.8), Inches(1.3), Inches(11.73), Inches(1.4), BG_CARD)
    card_tf = card.text_frame
    card_tf.word_wrap = True
    set_paragraph(card_tf, "If a company has your account, they know your name.",
                  font_size=24, color=YELLOW_ACCENT, bold=True, alignment=PP_ALIGN.CENTER,
                  space_after=Pt(10))
    add_paragraph(card_tf, "Mass phishing campaigns can't personalize every email. "
                  "Generic greetings are a strong sign of a phish.",
                  font_size=17, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    # Suspicious greetings (left)
    sus_card = add_rounded_rect(slide, Inches(0.8), Inches(3.2), Inches(5.6), Inches(3.3),
                                BG_EMAIL, RED_ACCENT)
    sus_tf = sus_card.text_frame
    sus_tf.word_wrap = True
    set_paragraph(sus_tf, "🚫  Suspicious Greetings", font_size=20, color=RED_ACCENT,
                  bold=True, alignment=PP_ALIGN.CENTER, space_after=Pt(16))
    suspicious = [
        "\"Dear Customer\"",
        "\"Dear User\"",
        "\"Dear Valued Member\"",
        "\"Dear Account Holder\"",
        "\"Hello Sir/Madam\"",
    ]
    for s in suspicious:
        add_paragraph(sus_tf, f"❌   {s}", font_size=18, color=LIGHT_GRAY, space_after=Pt(10))

    # Legitimate greetings (right)
    legit_card = add_rounded_rect(slide, Inches(6.93), Inches(3.2), Inches(5.6), Inches(3.3),
                                  BG_EMAIL, GREEN_ACCENT)
    legit_tf = legit_card.text_frame
    legit_tf.word_wrap = True
    set_paragraph(legit_tf, "✅  Legitimate Greetings", font_size=20, color=GREEN_ACCENT,
                  bold=True, alignment=PP_ALIGN.CENTER, space_after=Pt(16))
    legitimate = [
        "\"Hi John,\"",
        "\"Dear John Smith,\"",
        "\"Hello John,\"",
        "Uses your account username",
        "References your account number",
    ]
    for l in legitimate:
        add_paragraph(legit_tf, f"✅   {l}", font_size=18, color=LIGHT_GRAY, space_after=Pt(10))

    # Note
    tf = add_textbox(slide, Inches(1), Inches(6.8), Inches(11.33), Inches(0.5))
    set_paragraph(tf, "⚠️  Note: Spear phishing may use your real name. Generic greetings are "
                  "a red flag, but a personal greeting doesn't guarantee safety.",
                  font_size=14, color=DIM_GRAY, alignment=PP_ALIGN.CENTER, italic=True)

    add_speaker_notes(slide, (
        "If you have an account with a company, they know your name. If an email starts with "
        "'Dear Customer,' 'Dear User,' or 'Valued Member,' be skeptical. Mass phishing "
        "campaigns don't usually have the time or data to personalize every single email."
    ))


def build_slide_7(prs):
    """Transition to Game"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)
    add_rectangle(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.06), CYAN)

    # Game controller icon
    tf = add_textbox(slide, Inches(0), Inches(1.2), Inches(13.33), Inches(1.2))
    set_paragraph(tf, "🎮", font_size=70, alignment=PP_ALIGN.CENTER)

    # Title
    tf = add_textbox(slide, Inches(1.5), Inches(2.7), Inches(10.33), Inches(1.0))
    set_paragraph(tf, "LET'S PLAY:", font_size=28, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER,
                  space_after=Pt(4))
    add_paragraph(tf, "\"SPOT THE PHISH\"", font_size=46, color=CYAN, bold=True,
                  alignment=PP_ALIGN.CENTER)

    # Subtitle card
    card = add_rounded_rect(slide, Inches(2.5), Inches(4.4), Inches(8.33), Inches(1.6),
                            BG_CARD, CYAN)
    card_tf = card.text_frame
    card_tf.word_wrap = True
    set_paragraph(card_tf, "I'll show you a message.", font_size=20, color=WHITE,
                  alignment=PP_ALIGN.CENTER, space_after=Pt(8))
    add_paragraph(card_tf, "You tell me: Is it SAFE  ✅  or is it a SCAM  ❌  ?",
                  font_size=22, color=YELLOW_ACCENT, bold=True, alignment=PP_ALIGN.CENTER,
                  space_after=Pt(8))
    add_paragraph(card_tf, "Look for the red flags we just discussed!",
                  font_size=16, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    # Bottom
    tf = add_textbox(slide, Inches(3), Inches(6.6), Inches(7.33), Inches(0.5))
    set_paragraph(tf, "3 Scenarios  •  Interactive  •  Real-World Examples",
                  font_size=16, color=DIM_GRAY, alignment=PP_ALIGN.CENTER)

    add_rectangle(slide, Inches(0), Inches(7.44), Inches(13.33), Inches(0.06), CYAN)

    add_speaker_notes(slide, (
        "We've covered the theory. Now, let's see how well you can spot a scam in the wild. "
        "We are going to play a quick game of 'Spot the Phish.' I will show you a message, and "
        "I want you to tell me if it's safe, or if it's a scam."
    ))


def build_slide_8(prs):
    """Scenario 1 - The Fake Bank Alert (Question)"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)
    add_rectangle(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.06), YELLOW_ACCENT)

    # Scenario tag
    tag = add_rounded_rect(slide, Inches(0.6), Inches(0.3), Inches(2.5), Inches(0.45), YELLOW_ACCENT)
    tag_tf = tag.text_frame
    set_paragraph(tag_tf, "SCENARIO 1 OF 3", font_size=14, color=BG_DARK, bold=True,
                  alignment=PP_ALIGN.CENTER)

    tf = add_textbox(slide, Inches(3.4), Inches(0.25), Inches(9), Inches(0.7))
    set_paragraph(tf, "🏦  The Fake Bank Alert", font_size=32, color=WHITE, bold=True)

    # Email mockup
    email_card = add_rounded_rect(slide, Inches(1.8), Inches(1.2), Inches(9.73), Inches(5.0),
                                  BG_EMAIL, DIM_GRAY)
    email_tf = email_card.text_frame
    email_tf.word_wrap = True

    # Email header
    set_paragraph(email_tf, "From:    Chase Bank Security  <alerts@chase-secure-update12.com>",
                  font_size=14, color=LIGHT_GRAY, space_after=Pt(4), font_name="Consolas")
    add_paragraph(email_tf, "To:        you@email.com",
                  font_size=14, color=DIM_GRAY, space_after=Pt(4), font_name="Consolas")
    add_paragraph(email_tf, "Subject: URGENT: Fraudulent Activity Detected on Your Account",
                  font_size=14, color=RED_ACCENT, bold=True, space_after=Pt(2), font_name="Consolas")
    add_paragraph(email_tf, "─" * 72, font_size=10, color=DIM_GRAY, space_after=Pt(14))

    # Email body
    add_paragraph(email_tf, "Dear Customer,", font_size=17, color=WHITE, space_after=Pt(14))
    add_paragraph(email_tf,
                  "We have detected a login attempt from Russia on your Chase account. "
                  "For your security, please verify your identity immediately.",
                  font_size=16, color=LIGHT_GRAY, space_after=Pt(10))
    add_paragraph(email_tf,
                  "Click the link below within 12 hours to verify your identity, or your "
                  "account will be permanently locked.",
                  font_size=16, color=LIGHT_GRAY, space_after=Pt(14))

    # Button
    btn = add_rounded_rect(slide, Inches(4.5), Inches(4.95), Inches(4.33), Inches(0.55),
                           BUTTON_BLUE)
    btn_tf = btn.text_frame
    set_paragraph(btn_tf, "🔒  Verify Account Now", font_size=16, color=WHITE, bold=True,
                  alignment=PP_ALIGN.CENTER)

    # Hover tooltip
    tooltip = add_rounded_rect(slide, Inches(4.2), Inches(5.55), Inches(5.0), Inches(0.45),
                               RGBColor(0x33, 0x33, 0x33))
    tooltip_tf = tooltip.text_frame
    set_paragraph(tooltip_tf, "🔗 http://bit.ly/chase-auth-773", font_size=13,
                  color=YELLOW_ACCENT, font_name="Consolas", alignment=PP_ALIGN.CENTER)

    # Question prompt at bottom
    q_card = add_rounded_rect(slide, Inches(2.0), Inches(6.5), Inches(9.33), Inches(0.8),
                              BG_CARD, YELLOW_ACCENT)
    q_tf = q_card.text_frame
    q_tf.word_wrap = True
    set_paragraph(q_tf, "🤔  Is this email SAFE ✅ or a SCAM ❌ ?   What gives it away?",
                  font_size=20, color=YELLOW_ACCENT, bold=True, alignment=PP_ALIGN.CENTER)

    add_speaker_notes(slide, (
        "Take a look at this email. Raise your hand if you think this is a phish. "
        "What gives it away? Give the audience a moment to study the email and identify "
        "the red flags before moving to the next slide."
    ))


def build_slide_9(prs):
    """Scenario 1 - The Reveal"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)
    add_rectangle(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.06), RED_ACCENT)

    # Header
    tag = add_rounded_rect(slide, Inches(0.6), Inches(0.2), Inches(2.5), Inches(0.45), RED_ACCENT)
    tag_tf = tag.text_frame
    set_paragraph(tag_tf, "SCENARIO 1 — REVEAL", font_size=14, color=WHITE, bold=True,
                  alignment=PP_ALIGN.CENTER)

    tf = add_textbox(slide, Inches(3.4), Inches(0.15), Inches(9), Inches(0.7))
    set_paragraph(tf, "🚨  IT'S A PHISH!", font_size=34, color=RED_ACCENT, bold=True)

    # Email mockup (same as before but smaller/condensed)
    email_card = add_rounded_rect(slide, Inches(0.6), Inches(1.0), Inches(6.6), Inches(4.5),
                                  BG_EMAIL, RED_ACCENT)
    email_tf = email_card.text_frame
    email_tf.word_wrap = True

    set_paragraph(email_tf, "From: Chase Bank Security", font_size=13, color=LIGHT_GRAY,
                  font_name="Consolas", space_after=Pt(2))
    add_paragraph(email_tf, "  <alerts@chase-secure-update12.com>", font_size=12,
                  color=RED_ACCENT, font_name="Consolas", bold=True, space_after=Pt(4))
    add_paragraph(email_tf, "Subject: URGENT: Fraudulent Activity Detected",
                  font_size=13, color=LIGHT_GRAY, font_name="Consolas", space_after=Pt(2))
    add_paragraph(email_tf, "─" * 50, font_size=8, color=DIM_GRAY, space_after=Pt(8))
    add_paragraph(email_tf, "Dear Customer,", font_size=14, color=WHITE, space_after=Pt(8))
    add_paragraph(email_tf, "We have detected a login attempt from Russia...",
                  font_size=13, color=LIGHT_GRAY, space_after=Pt(6))
    add_paragraph(email_tf, "Click within 12 hours or your account will be permanently locked.",
                  font_size=13, color=LIGHT_GRAY, space_after=Pt(10))
    add_paragraph(email_tf, "   [ 🔒 Verify Account Now ]", font_size=14, color=CYAN,
                  bold=True, space_after=Pt(4))
    add_paragraph(email_tf, "   → http://bit.ly/chase-auth-773", font_size=11,
                  color=YELLOW_ACCENT, font_name="Consolas")

    # Red flag annotations on the right
    flags_x = Inches(7.6)
    flags_width = Inches(5.3)
    flag_data = [
        ("🔴  Fake Domain", "Not chase.com → chase-secure-update12.com\nA real bank uses its official domain.", Inches(1.0)),
        ("🔴  Generic Greeting", "\"Dear Customer\" — Chase would use\nyour actual name if you have an account.", Inches(2.6)),
        ("🔴  Urgency & Threats", "\"12 hours or permanently locked\" —\nDesigned to panic you into clicking.", Inches(4.0)),
        ("🔴  Suspicious Link", "Shortened bit.ly URL instead of a secure\nchase.com domain. Never trust short URLs.", Inches(5.3)),
    ]
    for title, desc, top in flag_data:
        flag_card = add_rounded_rect(slide, flags_x, top, flags_width, Inches(1.2),
                                     BG_CARD, RED_ACCENT)
        flag_tf = flag_card.text_frame
        flag_tf.word_wrap = True
        set_paragraph(flag_tf, title, font_size=16, color=RED_ACCENT, bold=True,
                      space_after=Pt(4))
        add_paragraph(flag_tf, desc, font_size=13, color=LIGHT_GRAY, space_after=Pt(2))

    # Bottom
    tip = add_rounded_rect(slide, Inches(1.5), Inches(6.8), Inches(10.33), Inches(0.6),
                           BG_CARD, GREEN_ACCENT)
    tip_tf = tip.text_frame
    tip_tf.word_wrap = True
    set_paragraph(tip_tf, "✅  Always go directly to chase.com or call the number on your card. "
                  "Never click links in suspicious emails.",
                  font_size=15, color=GREEN_ACCENT, alignment=PP_ALIGN.CENTER)

    add_speaker_notes(slide, (
        "Let's break this down. Red Flag 1: The sender address is not chase.com — it's "
        "chase-secure-update12.com. Red Flag 2: It says 'Dear Customer' instead of your "
        "name. Red Flag 3: The urgency — '12 hours or permanently locked.' And Red Flag 4: "
        "A shortened bit.ly link instead of a secure, official chase.com domain. If you ever "
        "get an email like this, go directly to the company's website by typing it yourself, "
        "or call the number on the back of your card."
    ))


def build_slide_10(prs):
    """Scenario 2 - CEO Fraud SMS (Question)"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)
    add_rectangle(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.06), YELLOW_ACCENT)

    tag = add_rounded_rect(slide, Inches(0.6), Inches(0.3), Inches(2.5), Inches(0.45), YELLOW_ACCENT)
    tag_tf = tag.text_frame
    set_paragraph(tag_tf, "SCENARIO 2 OF 3", font_size=14, color=BG_DARK, bold=True,
                  alignment=PP_ALIGN.CENTER)

    tf = add_textbox(slide, Inches(3.4), Inches(0.25), Inches(9), Inches(0.7))
    set_paragraph(tf, "💬  The CEO Fraud (SMS)", font_size=32, color=WHITE, bold=True)

    # Phone mockup frame
    phone_bg = add_rounded_rect(slide, Inches(3.8), Inches(1.2), Inches(5.73), Inches(5.2),
                                RGBColor(0x10, 0x10, 0x20), DIM_GRAY)

    # Phone header
    tf = add_textbox(slide, Inches(4.2), Inches(1.4), Inches(4.93), Inches(0.5))
    set_paragraph(tf, "📱  +1 (347) 555-0192", font_size=16, color=DIM_GRAY,
                  alignment=PP_ALIGN.CENTER, font_name="Consolas")

    tf = add_textbox(slide, Inches(4.2), Inches(1.85), Inches(4.93), Inches(0.35))
    set_paragraph(tf, "Unknown Number  •  Today 2:47 PM", font_size=12, color=DIM_GRAY,
                  alignment=PP_ALIGN.CENTER)

    # SMS bubble
    sms_bubble = add_rounded_rect(slide, Inches(4.0), Inches(2.5), Inches(5.33), Inches(3.2),
                                  SMS_BUBBLE, DIM_GRAY)
    sms_tf = sms_bubble.text_frame
    sms_tf.word_wrap = True
    set_paragraph(sms_tf, "Hi, it's David Chen.", font_size=16, color=WHITE,
                  space_after=Pt(10))
    add_paragraph(sms_tf, "I'm in a meeting right now and can't talk, but I need "
                  "a quick favor.", font_size=16, color=WHITE, space_after=Pt(10))
    add_paragraph(sms_tf, "Can you buy 5 × $100 Apple gift cards for a client presentation? "
                  "I will reimburse you by the end of the day.",
                  font_size=16, color=WHITE, space_after=Pt(10))
    add_paragraph(sms_tf, "Please send me the codes as soon as possible. Thanks!",
                  font_size=16, color=WHITE, space_after=Pt(4))

    # Question prompt
    q_card = add_rounded_rect(slide, Inches(2.0), Inches(6.7), Inches(9.33), Inches(0.7),
                              BG_CARD, YELLOW_ACCENT)
    q_tf = q_card.text_frame
    q_tf.word_wrap = True
    set_paragraph(q_tf, "🤔  Your \"boss\" just texted you. Do you buy the gift cards?",
                  font_size=20, color=YELLOW_ACCENT, bold=True, alignment=PP_ALIGN.CENTER)

    add_speaker_notes(slide, (
        "You get this text from your boss, Mr. David Chen, the CEO. He says he's in a meeting "
        "and can't talk, but he needs you to buy 5 $100 Apple gift cards and send him the codes. "
        "Do you buy the gift cards? Why or why not? Give the audience time to discuss."
    ))


def build_slide_11(prs):
    """Scenario 2 - The Reveal"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)
    add_rectangle(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.06), RED_ACCENT)

    tag = add_rounded_rect(slide, Inches(0.6), Inches(0.2), Inches(2.5), Inches(0.45), RED_ACCENT)
    tag_tf = tag.text_frame
    set_paragraph(tag_tf, "SCENARIO 2 — REVEAL", font_size=14, color=WHITE, bold=True,
                  alignment=PP_ALIGN.CENTER)

    tf = add_textbox(slide, Inches(3.4), Inches(0.15), Inches(9), Inches(0.7))
    set_paragraph(tf, "🚨  IT'S A SCAM!   Never buy gift cards for \"your boss\"",
                  font_size=30, color=RED_ACCENT, bold=True)

    # Condensed SMS on the left
    phone = add_rounded_rect(slide, Inches(0.6), Inches(1.1), Inches(5.0), Inches(4.2),
                             RGBColor(0x10, 0x10, 0x20), RED_ACCENT)
    phone_tf = phone.text_frame
    phone_tf.word_wrap = True
    set_paragraph(phone_tf, "📱  +1 (347) 555-0192  (Unknown)", font_size=13,
                  color=DIM_GRAY, font_name="Consolas", space_after=Pt(8))
    add_paragraph(phone_tf, "\"Hi, it's David Chen. I'm in a meeting and can't talk...\"",
                  font_size=14, color=LIGHT_GRAY, space_after=Pt(6))
    add_paragraph(phone_tf, "\"...buy 5 × $100 Apple gift cards...\"",
                  font_size=14, color=LIGHT_GRAY, space_after=Pt(6))
    add_paragraph(phone_tf, "\"...send me the codes as soon as possible.\"",
                  font_size=14, color=LIGHT_GRAY)

    # Red flags on the right
    flags_x = Inches(6.2)
    flags_width = Inches(6.5)
    flag_data = [
        ("🔴  Gift Cards = Scam Currency",
         "Gift cards are untraceable and unrecoverable.\nNo legitimate business transaction uses gift cards\nas payment. This is the #1 hallmark of a scam.",
         Inches(1.1)),
        ("🔴  \"I Can't Talk Right Now\"",
         "This prevents you from calling your boss to verify.\nThe scammer creates a reason you can only text,\nkeeping them in control of the conversation.",
         Inches(2.8)),
        ("🔴  Unknown Phone Number",
         "The number isn't your boss's saved contact.\nThe area code may not even match your company's\nlocation. Always verify through known channels.",
         Inches(4.5)),
    ]
    for title, desc, top in flag_data:
        flag_card = add_rounded_rect(slide, flags_x, top, flags_width, Inches(1.4),
                                     BG_CARD, RED_ACCENT)
        flag_tf = flag_card.text_frame
        flag_tf.word_wrap = True
        set_paragraph(flag_tf, title, font_size=17, color=RED_ACCENT, bold=True,
                      space_after=Pt(6))
        add_paragraph(flag_tf, desc, font_size=14, color=LIGHT_GRAY, space_after=Pt(2))

    # Bottom tip
    tip = add_rounded_rect(slide, Inches(1.5), Inches(6.4), Inches(10.33), Inches(0.9),
                           BG_CARD, GREEN_ACCENT)
    tip_tf = tip.text_frame
    tip_tf.word_wrap = True
    set_paragraph(tip_tf, "✅  ALWAYS verify unusual requests through a separate, known channel. "
                  "Call your boss on their saved number, walk to their office, or use company chat.",
                  font_size=15, color=GREEN_ACCENT, alignment=PP_ALIGN.CENTER)

    add_speaker_notes(slide, (
        "This is a classic CEO fraud / business email compromise scam. Three red flags: "
        "First, gift cards are the currency of scammers — they are untraceable and unrecoverable. "
        "No legitimate business transaction uses gift cards. Second, 'I'm in a meeting and can't "
        "talk' is designed to prevent you from calling your boss to verify. Third, it's from an "
        "unknown number, not your boss's saved contact. Always verify unusual requests through a "
        "separate, known communication channel."
    ))


def build_slide_12(prs):
    """Scenario 3 - Package Delivery (Question)"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)
    add_rectangle(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.06), YELLOW_ACCENT)

    tag = add_rounded_rect(slide, Inches(0.6), Inches(0.3), Inches(2.5), Inches(0.45), YELLOW_ACCENT)
    tag_tf = tag.text_frame
    set_paragraph(tag_tf, "SCENARIO 3 OF 3", font_size=14, color=BG_DARK, bold=True,
                  alignment=PP_ALIGN.CENTER)

    tf = add_textbox(slide, Inches(3.4), Inches(0.25), Inches(9), Inches(0.7))
    set_paragraph(tf, "📦  The Package Delivery", font_size=32, color=WHITE, bold=True)

    # Email mockup — this one looks very "authentic"
    email_card = add_rounded_rect(slide, Inches(1.8), Inches(1.2), Inches(9.73), Inches(5.0),
                                  BG_EMAIL, DIM_GRAY)
    email_tf = email_card.text_frame
    email_tf.word_wrap = True

    # Email header
    set_paragraph(email_tf, "From:    FedEx Notifications  <no-reply@fedex.com>",
                  font_size=14, color=LIGHT_GRAY, space_after=Pt(4), font_name="Consolas")
    add_paragraph(email_tf, "To:        john.smith@email.com",
                  font_size=14, color=DIM_GRAY, space_after=Pt(4), font_name="Consolas")
    add_paragraph(email_tf, "Subject: Missed Delivery Notification — Action Required",
                  font_size=14, color=WHITE, bold=True, space_after=Pt(2), font_name="Consolas")
    add_paragraph(email_tf, "─" * 72, font_size=10, color=DIM_GRAY, space_after=Pt(12))

    # FedEx branding line
    add_paragraph(email_tf, "FedEx  |  Shipping & Delivery Services", font_size=18,
                  color=ORANGE_ACCENT, bold=True, space_after=Pt(14))

    # Body
    add_paragraph(email_tf, "Hi John,", font_size=17, color=WHITE, space_after=Pt(10))
    add_paragraph(email_tf,
                  "We attempted to deliver your package today but no one was available "
                  "to sign for it. Please review the attached invoice to confirm your "
                  "delivery details and reschedule.",
                  font_size=16, color=LIGHT_GRAY, space_after=Pt(14))

    # Attachment mockup
    attach_box = add_rounded_rect(slide, Inches(2.5), Inches(4.9), Inches(5.0), Inches(0.6),
                                  RGBColor(0x33, 0x33, 0x44), ORANGE_ACCENT)
    attach_tf = attach_box.text_frame
    set_paragraph(attach_tf, "📎  Delivery_Invoice_Details.exe    (247 KB)",
                  font_size=15, color=ORANGE_ACCENT, bold=True, alignment=PP_ALIGN.CENTER)

    # Note: looks legit
    tf = add_textbox(slide, Inches(2.5), Inches(5.7), Inches(8.33), Inches(0.5))
    set_paragraph(tf, "⚠️  This email uses your real name and a real-looking sender address...",
                  font_size=15, color=YELLOW_ACCENT, italic=True, alignment=PP_ALIGN.CENTER)

    # Question prompt
    q_card = add_rounded_rect(slide, Inches(2.0), Inches(6.5), Inches(9.33), Inches(0.8),
                              BG_CARD, YELLOW_ACCENT)
    q_tf = q_card.text_frame
    q_tf.word_wrap = True
    set_paragraph(q_tf, "🤔  Looks pretty authentic, right?   Is it SAFE ✅ or a SCAM ❌ ?",
                  font_size=20, color=YELLOW_ACCENT, bold=True, alignment=PP_ALIGN.CENTER)

    add_speaker_notes(slide, (
        "This one is tricky. It uses your real name — 'Hi John' — and the sender address "
        "looks like it's from fedex.com. The branding looks professional. Is it safe? "
        "Take a close look at the attachment. Give the audience time to study it."
    ))


def build_slide_13(prs):
    """Scenario 3 - The Reveal"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)
    add_rectangle(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.06), RED_ACCENT)

    tag = add_rounded_rect(slide, Inches(0.6), Inches(0.2), Inches(2.5), Inches(0.45), RED_ACCENT)
    tag_tf = tag.text_frame
    set_paragraph(tag_tf, "SCENARIO 3 — REVEAL", font_size=14, color=WHITE, bold=True,
                  alignment=PP_ALIGN.CENTER)

    tf = add_textbox(slide, Inches(3.4), Inches(0.15), Inches(9), Inches(0.7))
    set_paragraph(tf, "🚨  IT'S A MALWARE TRAP!", font_size=34, color=RED_ACCENT, bold=True)

    # Condensed email on left
    email_card = add_rounded_rect(slide, Inches(0.6), Inches(1.0), Inches(5.5), Inches(4.0),
                                  BG_EMAIL, RED_ACCENT)
    email_tf = email_card.text_frame
    email_tf.word_wrap = True
    set_paragraph(email_tf, "From: FedEx Notifications", font_size=13, color=LIGHT_GRAY,
                  font_name="Consolas", space_after=Pt(2))
    add_paragraph(email_tf, "  <no-reply@fedex.com>", font_size=12, color=LIGHT_GRAY,
                  font_name="Consolas", space_after=Pt(4))
    add_paragraph(email_tf, "─" * 42, font_size=8, color=DIM_GRAY, space_after=Pt(8))
    add_paragraph(email_tf, "\"Hi John, we tried to deliver a package...\"",
                  font_size=13, color=LIGHT_GRAY, space_after=Pt(10))

    # Highlighted attachment
    add_paragraph(email_tf, "📎  Delivery_Invoice_Details .exe",
                  font_size=16, color=RED_ACCENT, bold=True, space_after=Pt(6))
    add_paragraph(email_tf, "          ↑  THIS IS MALWARE!",
                  font_size=15, color=RED_ACCENT, bold=True)

    # Red flags on the right
    flags_x = Inches(6.6)
    flags_width = Inches(6.1)

    # Flag 1 - The Attachment
    flag1 = add_rounded_rect(slide, flags_x, Inches(1.0), flags_width, Inches(2.2),
                             BG_CARD, RED_ACCENT)
    flag1_tf = flag1.text_frame
    flag1_tf.word_wrap = True
    set_paragraph(flag1_tf, "🔴  The Attachment Is Malware", font_size=18, color=RED_ACCENT,
                  bold=True, space_after=Pt(8))
    add_paragraph(flag1_tf, "The file ends in .exe — an executable program!",
                  font_size=15, color=WHITE, bold=True, space_after=Pt(6))
    add_paragraph(flag1_tf, "FedEx, UPS, and USPS never send:", font_size=14,
                  color=LIGHT_GRAY, space_after=Pt(4))
    add_paragraph(flag1_tf, "  ❌  .exe files    ❌  .zip archives    ❌  .scr files",
                  font_size=14, color=RED_ACCENT, bold=True, space_after=Pt(6))
    add_paragraph(flag1_tf, "They send tracking numbers you look up on their site.",
                  font_size=14, color=LIGHT_GRAY)

    # Flag 2 - Sender Spoofing
    flag2 = add_rounded_rect(slide, flags_x, Inches(3.5), flags_width, Inches(2.0),
                             BG_CARD, ORANGE_ACCENT)
    flag2_tf = flag2.text_frame
    flag2_tf.word_wrap = True
    set_paragraph(flag2_tf, "🟠  Sender Spoofing", font_size=18, color=ORANGE_ACCENT,
                  bold=True, space_after=Pt(8))
    add_paragraph(flag2_tf, "Even though it says no-reply@fedex.com, attackers can",
                  font_size=14, color=LIGHT_GRAY, space_after=Pt(2))
    add_paragraph(flag2_tf, "FAKE (spoof) the sender address to look legitimate.",
                  font_size=15, color=WHITE, bold=True, space_after=Pt(8))
    add_paragraph(flag2_tf, "This is why checking links and attachments is your\n"
                  "ultimate line of defense — the \"From\" field can lie.",
                  font_size=14, color=LIGHT_GRAY)

    # Key takeaway
    takeaway = add_rounded_rect(slide, Inches(0.8), Inches(5.9), Inches(11.73), Inches(1.4),
                                BG_CARD, GREEN_ACCENT)
    takeaway_tf = takeaway.text_frame
    takeaway_tf.word_wrap = True
    set_paragraph(takeaway_tf, "🛡️  KEY TAKEAWAY", font_size=18, color=GREEN_ACCENT, bold=True,
                  alignment=PP_ALIGN.CENTER, space_after=Pt(8))
    add_paragraph(takeaway_tf,
                  "Even when the sender looks real and the email uses your name — "
                  "ALWAYS check attachments and links. If you didn't request it, "
                  "don't open it. Verify through official channels.",
                  font_size=16, color=WHITE, alignment=PP_ALIGN.CENTER)

    add_speaker_notes(slide, (
        "Even though the sender says no-reply@fedex.com and uses your real name, the "
        "attachment gives it away. Delivery_Invoice_Details.exe is an executable file — "
        "that's malware. Delivery companies never send .exe, .zip, or .scr files. They "
        "give you a tracking number you can look up on their website.\n\n"
        "Also, explain that attackers can sometimes fake or 'spoof' the sender address "
        "to look real. This is why checking attachments and links is your ultimate defense. "
        "The 'From' field can lie.\n\n"
        "Thank you all for participating! Remember: when in doubt, verify through a separate "
        "channel. Never click, download, or respond in a panic. Stay safe!"
    ))


# ══════════════════════════════════════════════════════════════════════════
# MAIN
# ══════════════════════════════════════════════════════════════════════════

def main():
    prs = Presentation()
    # Set widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    build_slide_1(prs)
    build_slide_2(prs)
    build_slide_3(prs)
    build_slide_4(prs)
    build_slide_5(prs)
    build_slide_6(prs)
    build_slide_7(prs)
    build_slide_8(prs)
    build_slide_9(prs)
    build_slide_10(prs)
    build_slide_11(prs)
    build_slide_12(prs)
    build_slide_13(prs)

    output_path = os.path.join(os.path.dirname(__file__), "Phishing_Awareness_Training.pptx")
    prs.save(output_path)
    print(f"[OK] Presentation saved to: {output_path}")
    print(f"     {len(prs.slides)} slides generated")


if __name__ == "__main__":
    main()
